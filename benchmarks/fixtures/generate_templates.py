"""生成 M6 合成模板 fixture（§61：只用 synthetic / public / de-identified 材料）。

运行：python benchmarks/fixtures/generate_templates.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

HERE = Path(__file__).resolve().parent


def merge_education_template() -> Path:
    doc = Document()
    doc.add_heading("《融合教育》课程脚本模板", level=0)
    doc.add_paragraph("项目三 观察与记录（合成模板，仅用于 M6 benchmark）")
    doc.add_paragraph("第一节理论 13 分钟 | 第二节理论 13 分钟 | 第三节理论 13 分钟")
    for index in (1, 2, 3):
        doc.add_heading(f"任务{index}：【填写任务名称】", level=1)
        doc.add_paragraph("学习目标：【填写】")
        doc.add_paragraph("一、导入（2 分钟）")
        doc.add_paragraph("二、讲解要点（6 分钟）")
        doc.add_paragraph("三、案例（4 分钟）")
        doc.add_paragraph("四、小结（1 分钟）")
    doc.add_paragraph("称谓约束：称呼儿童请使用“儿童”或“幼儿”，不得使用标签化称谓。")
    path = HERE / "template_merge_education.docx"
    doc.save(path)
    return path


def childcare_template() -> Path:
    doc = Document()
    doc.add_heading("《托育机构管理实务》课程脚本模板", level=0)
    doc.add_paragraph("第八章 托育机构文化建设与品牌管理（合成模板，仅用于 M6 benchmark）")
    for section in ("8.1 办托理念与价值体系构建", "8.2 文化建设的策略", "8.3 员工行为规范与服务礼仪建设", "8.4 品牌形象与口碑传播管理"):
        doc.add_heading(section, level=1)
        doc.add_paragraph("本节字数：【1800–2100】")
        doc.add_paragraph("学习目标：【填写】")
        doc.add_paragraph("导入 → 讲解 → 案例 → 小结")
    doc.add_paragraph("称谓约束：统一使用“婴幼儿”。")
    path = HERE / "template_childcare_management.docx"
    doc.save(path)
    return path


def main() -> None:
    for path in (merge_education_template(), childcare_template()):
        print("wrote", path)


if __name__ == "__main__":
    main()
