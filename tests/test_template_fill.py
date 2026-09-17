from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from pebs import pptx_builder


def test_template_pptx_supplies_layouts(tmp_path):
    template = tmp_path / "brand.pptx"
    base = Presentation()
    base.save(str(template))

    output = tmp_path / "deck.pptx"
    built = pptx_builder.build_deck(
        course="测试课程",
        plan_rows=[
            {"title": "封面页", "layout_archetype": "cover", "family": "cover", "citation_labels": []},
            {"title": "内容页", "layout_archetype": "bullets", "family": "title_content", "points": ["要点一"], "citation_labels": []},
        ],
        output_path=output,
        template_path=template,
    )
    assert built["template_used"] == str(template)
    assert built["stats"]["slides"] == 2
    deck = Presentation(str(output))
    assert len(deck.slides) == 2
    first_text = "\n".join(
        shape.text_frame.text for shape in deck.slides[0].shapes if shape.has_text_frame
    )
    assert "封面页" in first_text


def test_without_template_uses_own_layout(tmp_path):
    output = tmp_path / "plain.pptx"
    built = pptx_builder.build_deck(
        course="测试课程",
        plan_rows=[{"title": "封面", "layout_archetype": "cover", "citation_labels": []}],
        output_path=output,
    )
    assert built["template_used"] == ""
    deck = Presentation(str(output))
    assert len(deck.slides) == 1
