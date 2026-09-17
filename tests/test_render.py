from __future__ import annotations

from pathlib import Path

import pytest
from conftest import REQUEST_1, run_build

from pebs import render
from pebs.engine import PlanEditRejected

pytestmark = pytest.mark.integration


def test_find_renderer_checks_common_install_paths(monkeypatch):
    monkeypatch.setattr(render.shutil, "which", lambda name: None)
    found = render.find_renderer()
    if not found:
        pytest.skip("no LibreOffice/PowerPoint renderer on this host")


def test_edit_slide_rebuilds_deck_and_keeps_other_artifacts(render_engine):
    engine = render_engine
    run_id, changeset_id = run_build(engine, REQUEST_1)
    engine.accept(changeset_id)
    script_rev = engine.store.accepted_rev_id("script:sec1")
    result = engine.edit_slide(2, {"title": "新的对比页标题", "points": ["要点A", "要点B"]})
    assert result["ok"] is True
    assert result["patched"] == ["points", "title"]

    slide_plan = engine.store.get_revision(engine.store.revisions_of("slide_plan")[-1])["content"]
    assert slide_plan["rows"][1]["title"] == "新的对比页标题"
    deck = engine.store.get_revision(engine.store.revisions_of("pptx_deck")[-1])["content"]
    assert deck["slides"][1]["title"] == "新的对比页标题"
    preview = engine.store.get_revision(engine.store.revisions_of("preview")[-1])["content"]["markdown"]
    assert "新的对比页标题" in preview

    accept = engine.accept(result["changeset_id"])
    assert accept["status"] == "accepted"
    assert engine.store.accepted_rev_id("script:sec1") == script_rev

    from pptx import Presentation

    deck_content = engine.store.accepted_content("pptx_deck")
    presentation = Presentation(deck_content["path"])
    texts = "\n".join(
        shape.text_frame.text for shape in presentation.slides[1].shapes if shape.has_text_frame
    )
    assert "新的对比页标题" in texts


def test_edit_slide_rejects_unknown_fields_and_bad_points(render_engine):
    engine = render_engine
    run_id, changeset_id = run_build(engine, REQUEST_1)
    engine.accept(changeset_id)
    with pytest.raises(PlanEditRejected):
        engine.edit_slide(2, {"diagram_id": "d9"})
    with pytest.raises(PlanEditRejected):
        engine.edit_slide(2, {"points": ["1", "2", "3", "4", "5", "6", "7"]})
    with pytest.raises(PlanEditRejected):
        engine.edit_slide(999, {"title": "x"})


def test_renderer_thumbnails_are_recorded(render_engine):
    engine = render_engine
    run_id, changeset_id = run_build(engine, REQUEST_1)
    deck = engine.store.get_revision(engine.store.revisions_of("pptx_deck")[-1])["content"]
    assert deck["renderer"]
    assert deck["thumbnails"], "渲染器可用时必须生成缩略图"
    assert deck["qa"]["status"] == "PASS"
    assert len(deck["thumbnails"]) == len(deck["slides"])
    from pebs import pptx_foundry

    if pptx_foundry.available():
        assert deck["delivery_check"], "已登记 ppt-foundry 时交付检查应执行"
    else:
        assert deck["delivery_check"] == {}
    renderer_check = next(check for check in deck["qa"]["checks"] if check["check"] == "真实渲染")
    assert renderer_check["status"] == "PASS"


def test_g8_passes_with_renderer(render_engine):
    engine = render_engine
    run_id, changeset_id = run_build(engine, REQUEST_1)
    engine.accept(changeset_id)
    g8 = engine.store.accepted_content("gate:G8:export_manifest")
    assert g8["status"] == "PASS", g8["issues"]
