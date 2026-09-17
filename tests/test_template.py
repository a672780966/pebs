from __future__ import annotations

import docx
import pytest
from conftest import REQUEST_1, run_build

from pebs import gates
from pebs.template_parse import parse_template


def test_markdown_template_parsing(tmp_path):
    path = tmp_path / "template.md"
    path.write_text(
        "# 课程脚本模板\n\n称谓统一使用「婴幼儿」。\n\n| 教师讲解 | 画面提示 | 案例 |\n|---|---|---|\n|  |  |  |\n",
        encoding="utf-8",
    )
    spec = parse_template(path)
    assert spec["kind"] == "markdown"
    assert [c["name"] for c in spec["columns"]] == ["教师讲解", "画面提示", "案例"]
    assert "婴幼儿" in spec["terminology_rules"]


def test_docx_template_parsing(tmp_path):
    path = tmp_path / "template.docx"
    document = docx.Document()
    document.add_heading("课程脚本模板", level=1)
    document.add_paragraph("称谓统一使用「婴幼儿」。")
    table = document.add_table(rows=2, cols=3)
    for idx, name in enumerate(["教师讲解", "画面提示", "案例"]):
        table.rows[0].cells[idx].text = name
    document.save(str(path))
    spec = parse_template(path)
    assert spec["kind"] == "docx_table"
    assert [c["name"] for c in spec["columns"]] == ["教师讲解", "画面提示", "案例"]
    assert "婴幼儿" in spec["terminology_rules"]


def test_g7_fails_when_template_terminology_not_followed(engine, tmp_path):
    path = tmp_path / "template.md"
    path.write_text(
        "# 模板\n\n称谓统一使用「婴幼儿」。\n\n| 教师讲解 | 画面提示 | 案例 |\n|---|---|---|\n|  |  |  |\n",
        encoding="utf-8",
    )
    run_id, changeset_id = run_build(engine, REQUEST_1, template_path=path)
    engine.accept(changeset_id)
    ctx = gates.GateContext(store=engine.store, evidence=engine.evidence)
    result = gates.g7_template(ctx, "script:sec1")
    assert result["status"] == "FAIL"
    assert any("婴幼儿" in issue["reason"] for issue in result["issues"])


def test_g7_not_applicable_without_template(engine):
    run_id, changeset_id = run_build(engine, REQUEST_1)
    engine.accept(changeset_id)
    ctx = gates.GateContext(store=engine.store, evidence=engine.evidence)
    result = gates.g7_template(ctx, "script:sec1")
    assert result["status"] == "NOT_APPLICABLE"
    assert result["not_applicable_reason"]
