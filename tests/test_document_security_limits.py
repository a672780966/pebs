"""M6 §59 Security Track：上传/导入的文档级与解压级限制。

Stage C 已覆盖文件数/单文件体积/总量；这里补齐 DOCX 正文字数、PDF 页数、
模板表格数、解压炸弹（zip 压缩比/解压总量）与损坏文档的显式拒绝。
所有拒绝都必须**显式报错**，不允许静默截断或跳过。
"""

from __future__ import annotations

import zipfile

import pytest

from pebs import config, template_parse, server as server_mod
from pebs.template_parse import ImportLimitExceeded, ParseError


@pytest.fixture
def client(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient

    monkeypatch.setattr(config, "PROJECTS_DIR", tmp_path / "projects")
    monkeypatch.setattr(
        config,
        "PROVIDERS",
        {"llm": {"kind": "openai_compatible", "enabled": False}, "research": {"enabled": False}},
    )
    server_mod._engines.clear()
    return TestClient(server_mod.app)


def test_decompression_bomb_is_rejected(tmp_path, monkeypatch):
    """§59：解压总量越界必须拒绝（用可复现的阈值注入，不构造真实炸弹）。"""
    from docx import Document

    path = tmp_path / "big_ratio.docx"
    document = Document()
    document.add_paragraph("重复内容" * 200)
    document.save(path)

    monkeypatch.setattr(template_parse, "MAX_DECOMPRESSED_MIB", 0)  # 任何解压体积都越界
    with pytest.raises(ImportLimitExceeded) as excinfo:
        template_parse.check_archive_safety(path)
    assert excinfo.value.http_status == 413
    assert "解压体积异常" in str(excinfo.value)

    monkeypatch.setattr(template_parse, "MAX_DECOMPRESSED_MIB", 200)
    monkeypatch.setattr(template_parse, "MAX_DECOMPRESSION_RATIO", 1)  # 压缩比 > 1 即越界
    with pytest.raises(ImportLimitExceeded):
        template_parse.check_archive_safety(path)


def test_normal_docx_passes_archive_check(tmp_path):
    from docx import Document

    path = tmp_path / "ok.docx"
    document = Document()
    document.add_paragraph("正常模板")
    document.save(path)
    template_parse.check_archive_safety(path)


def test_malformed_docx_is_rejected_with_clear_error(tmp_path):
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a zip at all")
    with pytest.raises(ParseError) as excinfo:
        template_parse.check_archive_safety(broken)
    assert "损坏" in str(excinfo.value)


def test_docx_char_limit_rejects_instead_of_truncating(tmp_path, monkeypatch):
    from docx import Document

    path = tmp_path / "big.docx"
    document = Document()
    document.add_paragraph("观察记录" * 50)
    document.save(path)
    rules = dict(config.RULES)
    limits = dict(rules.get("limits", {}))
    limits["max_docx_chars"] = 10
    rules["limits"] = limits
    monkeypatch.setattr(config, "RULES", rules)
    monkeypatch.setattr(template_parse, "MAX_DOCX_CHARS", 10)
    with pytest.raises(ImportLimitExceeded) as excinfo:
        template_parse.extract_text(path)
    assert "正文超过" in str(excinfo.value)
    assert "拆分" in str(excinfo.value)


def test_template_table_limit_rejects(tmp_path, monkeypatch):
    from docx import Document

    path = tmp_path / "many_tables.docx"
    document = Document()
    for _ in range(3):
        document.add_table(rows=1, cols=2)
    document.save(path)
    monkeypatch.setattr(template_parse, "MAX_TEMPLATE_TABLES", 1)
    with pytest.raises(ImportLimitExceeded) as excinfo:
        template_parse.parse_template(path)
    assert "表格数量超过" in str(excinfo.value)


def test_upload_filename_traversal_is_neutralised(client, tmp_path):
    """§59：上传文件名不得逃出项目 inputs 目录（只取 basename）。"""
    client.post("/api/projects", json={"project_id": "travtest"})
    res = client.post(
        "/api/projects/travtest/upload",
        files={"file": ("../../evil.docx", b"hello", "application/octet-stream")},
    )
    assert res.status_code == 200, res.text
    stored = res.json()["path"].replace("\\", "/")
    assert stored.endswith("inputs/evil.docx"), stored
    assert ".." not in stored
    target = (tmp_path / "projects" / "travtest" / "inputs" / "evil.docx").resolve()
    assert target.exists()


def test_malformed_document_surfaces_a_clear_error(client):
    """§59：损坏文档必须给出可解释错误，而不是 500/静默跳过。"""
    client.post("/api/projects", json={"project_id": "brokendoc"})
    upload = client.post(
        "/api/projects/brokendoc/upload",
        files={"file": ("broken.docx", b"definitely not a zip", "application/octet-stream")},
    )
    assert upload.status_code == 200
    res = client.post(
        "/api/projects/brokendoc/build",
        json={
            "request": "读取材料并写一节课程脚本，每节10到999字，每节至少1个案例。",
            "material_paths": [upload.json()["path"]],
        },
    )
    assert res.status_code in (200, 409)
    if res.status_code != 200:
        return
    import time

    run_id = res.json()["run_id"]
    deadline = time.time() + 60
    status = client.get(f"/api/projects/brokendoc/runs/{run_id}").json()
    while status["run"]["status"] == "running" and time.time() < deadline:
        time.sleep(0.3)
        status = client.get(f"/api/projects/brokendoc/runs/{run_id}").json()
    parse_step = next(step for step in status["steps"] if step["step_id"] == "parse_inputs")
    assert parse_step["status"] in ("FAILED", "BLOCKED")
    error = parse_step["error"] or ""
    assert "损坏" in error or "无效" in error or "无法" in error


def test_pdf_page_limit_rejects(tmp_path, monkeypatch):
    """§59：PDF 页数上限。"""
    pytest.importorskip("pypdf")
    from pypdf import PdfWriter

    path = tmp_path / "pages.pdf"
    writer = PdfWriter()
    for _ in range(3):
        writer.add_blank_page(width=200, height=200)
    with path.open("wb") as handle:
        writer.write(handle)
    monkeypatch.setattr(template_parse, "MAX_PDF_PAGES", 1)
    with pytest.raises(ImportLimitExceeded) as excinfo:
        template_parse.extract_text(path)
    assert "页数超过" in str(excinfo.value)
