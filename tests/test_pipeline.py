from __future__ import annotations

import json
from pathlib import Path

from conftest import REQUEST_1, REQUEST_2, FakeLLM, FakeResearch, MissingLLM, run_build

from pebs import gates

def _step_statuses(engine, run_id):
    return {s["step_id"]: s["status"] for s in engine.store.get_steps(run_id)}


def test_full_build_produces_formal_ready_artifacts(engine):
    run_id, changeset_id = run_build(engine, REQUEST_1)
    statuses = _step_statuses(engine, run_id)
    assert statuses["parse_inputs"] == "SUCCEEDED"
    assert statuses["learning_design"] == "SUCCEEDED"
    assert statuses["scripts"] == "SUCCEEDED"
    assert statuses["gates"] == "SUCCEEDED"
    assert statuses["export"] == "SUCCEEDED"
    assert engine.store.get_run(run_id)["status"] == "succeeded"

    claims = engine.evidence.claims()
    assert claims and all(c["status"] == "SUPPORTED" for c in claims)

    candidate_script = engine.store.revisions_of("script:sec1")[-1]
    assert engine.store.get_revision(candidate_script)["content"]["word_count"]["in_range"] is True

    result = engine.accept(changeset_id)
    assert result["status"] == "accepted"
    readiness = gates.export_readiness(engine.store)
    assert readiness["ready"] is True, readiness["blocking"]


def test_failed_provider_call_is_still_counted_as_consumption(engine):
    """§35 Cost / Quota Consumption：请求已发出但 provider 报错（如配额打满）也要记账。

    否则一次真实的配额耗尽会被记成 0 次调用，成本口径失真。
    """
    from pebs.providers import ProviderError

    class FlakyLLM(FakeLLM):
        def generate_json(self, *, task, system, prompt):
            self._bump(task)
            if task == "learning_design":
                raise ProviderError("codex exec 失败：usage limit")
            return super().generate_json(task=task, system=system, prompt=prompt)

    engine.llm = FlakyLLM()
    run_id, _ = run_build(engine, REQUEST_1)
    step = next(s for s in engine.store.get_steps(run_id) if s["step_id"] == "learning_design")
    assert step["status"] == "FAILED"
    assert int(step["model_calls"]) == 1, "失败的调用也要记到该步骤"
    assert int(engine.store.get_run(run_id)["calls_used"]) >= 1


def test_malformed_json_response_is_retried_once_then_fails(engine):
    """A14：畸形 JSON（长输出被截断）应有界重试一次；再失败则如实失败，不伪造内容。"""
    from pebs.providers import ProviderError

    class TruncatingLLM(FakeLLM):
        def __init__(self, fail_times: int):
            super().__init__()
            self.remaining = fail_times

        def generate_json(self, *, task, system, prompt):
            self._bump(task)
            if task == "storyboard" and self.remaining > 0:
                self.remaining -= 1
                raise ProviderError("provider JSON parse error: Unterminated string")
            return super().generate_json(task=task, system=system, prompt=prompt)

    # 第一次畸形、第二次正常 → 运行继续，且两次调用都被记账
    engine.llm = TruncatingLLM(fail_times=1)
    run_id, _ = run_build(engine, REQUEST_1)
    steps = {step["step_id"]: step for step in engine.store.get_steps(run_id)}
    assert steps["storyboard"]["status"] == "SUCCEEDED", steps["storyboard"]["error"]
    assert engine.store.get_run(run_id)["calls_used"] >= 2

    # 连续畸形 → 有界重试后失败，且没有产出伪造成品
    engine2 = type(engine)(engine.project_id + "-retry2", base_dir=engine.base.parent / "proj2")
    engine2.llm = TruncatingLLM(fail_times=99)
    engine2.research = engine.research
    run_id2, _ = run_build(engine2, REQUEST_1)
    steps2 = {step["step_id"]: step for step in engine2.store.get_steps(run_id2)}
    assert steps2["storyboard"]["status"] == "FAILED"
    assert "JSON parse error" in (steps2["storyboard"]["error"] or "")
    engine2.close()


def test_offline_without_provider_reports_unavailable_and_never_fabricates(engine):
    engine.llm = MissingLLM()
    run_id, changeset_id = run_build(engine, REQUEST_1)
    statuses = _step_statuses(engine, run_id)
    assert statuses["parse_inputs"] == "SUCCEEDED"
    assert statuses["requirements"] == "SUCCEEDED"
    assert statuses["learning_design"] == "BLOCKED"
    assert statuses["scripts"] == "BLOCKED"
    assert engine.store.get_run(run_id)["status"] == "blocked"
    step = next(s for s in engine.store.get_steps(run_id) if s["step_id"] == "learning_design")
    assert "Provider" in (step["error"] or "")
    assert engine.store.revisions_of("script:sec1") == []
    assert engine.store.revisions_of("claims") == []


def test_fixture_environment_is_recorded_in_runs_and_gate_results(engine):
    run_id, changeset_id = run_build(engine, REQUEST_1, environment="test_fixture")
    run = engine.store.get_run(run_id)
    assert run["environment"] == "test_fixture"
    gate_artifacts = [a for a in engine.store.list_artifacts() if a["artifact_type"] == "gate_result"]
    assert gate_artifacts
    for artifact in gate_artifacts:
        revision = engine.store.get_revision(artifact["accepted_rev"] or engine.store.revisions_of(artifact["artifact_id"])[-1])
        assert revision["fixture"] == 1
        assert revision["content"]["run_environment"] == "test_fixture"


def test_export_draft_label_when_gates_fail(engine):
    engine.research = FakeResearch(available=False)
    run_id, changeset_id = run_build(engine, REQUEST_1)
    engine.accept(changeset_id)
    readiness = gates.export_readiness(engine.store)
    assert readiness["ready"] is False
    result = engine.export_now("draft")
    assert result["ok"] is True
    files = result["manifest"]["files"]
    assert files
    labeled_roles = {"lesson_script", None}
    for entry in files:
        assert entry["label"] == "草稿—未通过 QA"
        role = entry.get("role", "lesson_script")
        if entry["kind"] == "docx" and role in labeled_roles:
            import docx

            document = docx.Document(entry["path"])
            text = "\n".join(p.text for p in document.paragraphs)
            assert "草稿—未通过 QA" in text
        elif entry["kind"] == "markdown" and role in labeled_roles:
            assert "草稿—未通过 QA" in Path(entry["path"]).read_text(encoding="utf-8")
        else:
            assert Path(entry["path"]).exists()


def test_formal_export_blocked_when_gates_fail(engine):
    engine.research = FakeResearch(available=False)
    run_id, changeset_id = run_build(engine, REQUEST_1)
    engine.accept(changeset_id)
    result = engine.export_now("formal")
    assert result["ok"] is False
    assert "正式导出被阻止" in result["reason"]


def test_docx_export_publishes_only_after_accept(engine, no_pptx):
    run_id, changeset_id = run_build(engine, REQUEST_1)
    engine.accept(changeset_id)
    formal_dir = Path(engine.store.base_dir) / "outputs" / "formal"
    before = sorted(p.name for p in formal_dir.iterdir()) if formal_dir.exists() else []

    result = engine.export_now("formal")
    assert result["ok"] is True
    staged_paths = [f["path"] for f in result["manifest"]["files"]]
    assert staged_paths and all("staging" in path for path in staged_paths)
    during = sorted(p.name for p in formal_dir.iterdir()) if formal_dir.exists() else []
    assert during == before

    accept = engine.accept(result["changeset_id"])
    assert accept["status"] == "accepted"
    published = {p.name for p in formal_dir.iterdir()}
    assert {Path(p).name for p in staged_paths} <= published
    import docx

    docx_file = next(f for f in result["manifest"]["files"] if f["kind"] == "docx")
    document = docx.Document(str(formal_dir / Path(docx_file["path"]).name))
    assert any("测试节" in p.text for p in document.paragraphs)
    assert not (Path(engine.store.base_dir) / "outputs" / "staging" / result["changeset_id"]).exists()
    script = engine.store.accepted_content("script:sec1")
    assert script["word_count"]["count"] > 0


def test_rejecting_formal_export_discards_staging(engine):
    run_id, changeset_id = run_build(engine, REQUEST_1)
    engine.accept(changeset_id)
    result = engine.export_now("formal")
    reject = engine.reject(result["changeset_id"])
    assert reject["status"] == "rejected"
    staging = Path(engine.store.base_dir) / "outputs" / "staging" / result["changeset_id"]
    assert not staging.exists()
