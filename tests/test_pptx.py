from __future__ import annotations

from pathlib import Path

from conftest import REQUEST_1, run_build

from pebs import gates, pptx_builder


def test_build_produces_editable_pptx(engine):
    run_id, changeset_id = run_build(engine, REQUEST_1)
    deck = engine.store.get_revision(engine.store.revisions_of("pptx_deck")[-1])["content"]
    path = Path(deck["path"])
    assert path.exists()
    from pptx import Presentation

    presentation = Presentation(str(path))
    assert len(presentation.slides) == 4
    assert deck["stats"]["tables"] >= 1
    assert deck["stats"]["shapes"] >= 4
    assert deck["stats"]["connectors"] >= 1
    assert deck["qa"]["jargon_hits"] == []
    assert deck["qa"]["status"] == "NEEDS_REVIEW"
    titles = [slide["title"] for slide in deck["slides"]]
    assert titles[0] == "观察记录：事实与判断"
    assert titles[-1] == "小结"
    second = presentation.slides[1]
    assert second.notes_slide.notes_text_frame.text.strip()

    engine.accept(changeset_id)
    g8 = engine.store.accepted_content("gate:G8:export_manifest")
    assert g8["status"] == "NEEDS_REVIEW"
    assert any("PPTX" in issue["reason"] for issue in g8["issues"])


def test_slide_plan_strips_unsupported_claims_and_keeps_valid_diagram(engine):
    run_id, changeset_id = run_build(engine, REQUEST_1)
    plan = engine.store.get_revision(engine.store.revisions_of("slide_plan")[-1])["content"]
    table_row = next(row for row in plan["rows"] if row["layout_archetype"] == "table")
    assert table_row["claim_refs"] == []
    diagram_row = next(row for row in plan["rows"] if row["layout_archetype"] == "diagram")
    assert diagram_row["diagram_id"] == "d1"
    step = next(s for s in engine.store.get_steps(run_id) if s["step_id"] == "slide_plan")
    assert "移除未获支持的引用" in (step["note"] or "")


def test_manifest_includes_presentation_and_downgrades_without_renderer(engine):
    run_id, changeset_id = run_build(engine, REQUEST_1)
    engine.accept(changeset_id)
    result = engine.export_now("formal")
    assert result["ok"] is True
    assert result["mode"] == "draft"
    roles = {entry.get("role") for entry in result["manifest"]["files"]}
    kinds = {entry["kind"] for entry in result["manifest"]["files"]}
    assert "presentation" in roles
    assert "pptx" in kinds


def test_formal_export_without_pptx_passes(engine, no_pptx):
    run_id, changeset_id = run_build(engine, REQUEST_1)
    engine.accept(changeset_id)
    assert gates.export_readiness(engine.store)["ready"] is True
    result = engine.export_now("formal")
    assert result["mode"] == "formal"
    assert result["g8"] == "PASS"
    assert all(entry["kind"] != "pptx" for entry in result["manifest"]["files"])


def test_jargon_scan_flags_internal_terms(tmp_path):
    output = tmp_path / "jargon.pptx"
    pptx_builder.build_deck(
        course="测试",
        plan_rows=[
            {
                "title": "内部术语页",
                "layout_archetype": "bullets",
                "points": ["参见 artifact 与 Animation Gate"],
                "citation_labels": [],
            }
        ],
        output_path=output,
    )
    inspection = pptx_builder.inspect_deck(output)
    assert inspection["jargon_hits"]
    assert {hit["token"] for hit in inspection["jargon_hits"]} >= {"artifact", "Animation Gate"}


def test_diagram_slide_uses_native_shapes(engine):
    run_id, changeset_id = run_build(engine, REQUEST_1)
    from pptx import Presentation

    deck = engine.store.get_revision(engine.store.revisions_of("pptx_deck")[-1])["content"]
    presentation = Presentation(deck["path"])
    diagram_slide = presentation.slides[1]
    kinds = [shape.shape_type for shape in diagram_slide.shapes]
    assert any(shape.has_text_frame for shape in diagram_slide.shapes)
    assert all(shape.shape_type != 13 for shape in diagram_slide.shapes)
    assert kinds
