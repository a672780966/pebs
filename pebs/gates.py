from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Any, Callable

from . import config, policies, router
from .evidence import EvidenceStore
from .store import Store, StoreError, content_hash, now_iso

GATE_ORDER = ["G1", "G2", "G3", "G4", "G5", "G6", "G7"]
ALL_GATES = GATE_ORDER + ["G8"]
PASSING = {"PASS", "NOT_APPLICABLE"}
LABEL_CHECK_ROLES = {"lesson_script", "lesson_plan", "worksheet", "worksheet_teacher"}
COLUMN_KIND_MAP = {
    "narration": ["讲解", "正文", "讲稿", "教师", "台词", "内容"],
    "visual": ["画面", "视觉", "媒体", "视觉提示"],
    "case": ["案例", "示例"],
    "interaction": ["互动", "提问", "活动", "讨论"],
}


def _norm(text: str) -> str:
    return re.sub(r"\s+", "", text or "")


@dataclass
class GateContext:
    store: Store
    evidence: EvidenceStore
    run_id: str = ""
    environment: str = "production"
    overrides: dict[str, str] = field(default_factory=dict)
    extras: dict[str, Any] = field(default_factory=dict)

    def content(self, artifact_id: str) -> Any | None:
        rev_id = self.overrides.get(artifact_id)
        if rev_id:
            return self.store.get_revision(rev_id)["content"]
        return self.store.accepted_content(artifact_id)

    def accepted(self, artifact_id: str) -> Any | None:
        return self.content(artifact_id)

    def accepted_rev(self, artifact_id: str) -> str | None:
        return self.overrides.get(artifact_id) or self.store.accepted_rev_id(artifact_id)


def gate_artifact_id(gate_id: str, target_artifact: str) -> str:
    return f"gate:{gate_id}:{target_artifact}"


def make_result(
    gate_id: str,
    *,
    target_artifact: str,
    target_rev_id: str,
    target_hash: str,
    status: str,
    issues: list[dict[str, str]] | None = None,
    dep_versions: dict[str, str] | None = None,
    run_id: str = "",
    environment: str = "production",
    na_reason: str | None = None,
    renderer: str | None = None,
) -> dict[str, Any]:
    result: dict[str, Any] = {
        "gate_id": gate_id,
        "check_version": config.CHECK_VERSIONS[gate_id],
        "target": {"artifact_id": target_artifact, "revision_id": target_rev_id, "content_hash": target_hash},
        "dep_versions": dep_versions or {},
        "rules_version": config.RULES_VERSION,
        "status": status,
        "issues": issues or [],
        "run_id": run_id,
        "run_environment": environment,
        "created_at": now_iso(),
    }
    if na_reason is not None:
        result["not_applicable_reason"] = na_reason
    if renderer is not None:
        result["renderer"] = renderer
    return result


class GateSetupError(Exception):
    pass


def _rev_info(ctx: GateContext, artifact_id: str) -> tuple[str, str] | None:
    rev_id = ctx.accepted_rev(artifact_id)
    if not rev_id:
        return None
    rev = ctx.store.get_revision(rev_id)
    return rev_id, rev["content_hash"]


def _require_rev(ctx: GateContext, artifact_id: str) -> tuple[str, str]:
    info = _rev_info(ctx, artifact_id)
    if info is None:
        raise GateSetupError(f"artifact has no accepted revision: {artifact_id}")
    return info


def _units_text(script: dict[str, Any]) -> list[tuple[str, str]]:
    return [(u.get("unit_id", ""), u.get("text", "")) for u in script.get("units", [])]


def _all_script_text(script: dict[str, Any]) -> str:
    return "\n".join(u.get("text", "") for u in script.get("units", []))


def _script_units_of_kind(script: dict[str, Any], kind: str) -> list[dict[str, Any]]:
    return [u for u in script.get("units", []) if u.get("kind") == kind]


def g1_requirements(ctx: GateContext, script_artifact: str) -> dict[str, Any]:
    rev_id, rev_hash = _require_rev(ctx, script_artifact)
    script = ctx.accepted(script_artifact) or {}
    reqs = ctx.accepted("requirements") or {}
    issues: list[dict[str, str]] = []
    for item in reqs.get("items", []):
        if item.get("status") == "conflicted":
            issues.append(
                {
                    "location": item.get("req_id", "requirements"),
                    "reason": f"存在未解决的要求冲突：{item.get('text', '')}",
                    "next_step": "在工作台中解决冲突后重跑",
                }
            )
    rules = reqs.get("word_rules", {})
    wc = script.get("word_count", {})
    if rules.get("min") and rules.get("max"):
        if wc.get("method") != config.RULES["word_counting"]["method"]:
            issues.append(
                {
                    "location": f"{script_artifact}:word_count",
                    "reason": f"字数统计口径为 {wc.get('method')}，与当前规则版本不一致",
                    "next_step": "按 config/rules.yaml 重新统计",
                }
            )
        elif not wc.get("in_range", False):
            issues.append(
                {
                    "location": f"{script_artifact}:word_count",
                    "reason": f"字数 {wc.get('count')} 不在要求范围 {rules.get('min')}–{rules.get('max')}",
                    "next_step": "扩写或压缩教师讲解正文后重跑",
                }
            )
    section = next(
        (s for s in reqs.get("sections", []) if s.get("section_id") == script.get("section_id")), None
    )
    if section:
        required_cases = section.get("required_case_count") or (
            1 if "case" in section.get("required_components", []) else 0
        )
        case_units = _script_units_of_kind(script, "case")
        if len(case_units) < required_cases:
            issues.append(
                {
                    "location": f"{script_artifact}:units",
                    "reason": f"要求每节至少 {required_cases} 个案例，脚本中实际 {len(case_units)} 个",
                    "next_step": "补充案例单元",
                }
            )
        if "assessment" in section.get("required_components", []):
            if ctx.accepted(f"assessment:{script.get('section_id')}") is None:
                issues.append(
                    {
                        "location": f"assessment:{script.get('section_id')}",
                        "reason": "要求包含形成性评价，未找到对应评价产物",
                        "next_step": "运行评价设计步骤",
                    }
                )
        outputs = reqs.get("outputs", [])
        if "lesson_plan" in outputs and ctx.accepted(f"lesson_plan:{script.get('section_id')}") is None:
            issues.append(
                {
                    "location": f"lesson_plan:{script.get('section_id')}",
                    "reason": "要求输出教案，未找到对应教案产物",
                    "next_step": "运行教案步骤",
                }
            )
        if "worksheet" in outputs and ctx.accepted(f"worksheet:{script.get('section_id')}") is None:
            issues.append(
                {
                    "location": f"worksheet:{script.get('section_id')}",
                    "reason": "要求输出学习单，未找到对应学习单产物",
                    "next_step": "运行学习单步骤",
                }
            )
    text = _all_script_text(script)
    for term in reqs.get("terminology", []):
        if term and term not in text:
            issues.append(
                {
                    "location": f"{script_artifact}:terminology",
                    "reason": f"术语要求未满足：文本中未出现“{term}”",
                    "next_step": f"按术语规则使用“{term}”改写",
                }
            )
    status = "FAIL" if issues else "PASS"
    deps = {k: v for k, v in {"requirements": ctx.accepted_rev("requirements")}.items() if v}
    return make_result(
        "G1",
        target_artifact=script_artifact,
        target_rev_id=rev_id,
        target_hash=rev_hash,
        status=status,
        issues=issues,
        dep_versions=deps,
        run_id=ctx.run_id,
        environment=ctx.environment,
    )


def g2_evidence(ctx: GateContext, script_artifact: str) -> dict[str, Any]:
    rev_id, rev_hash = _require_rev(ctx, script_artifact)
    script = ctx.accepted(script_artifact) or {}
    issues: list[dict[str, str]] = []
    ref_count = 0
    for unit in script.get("units", []):
        if unit.get("placeholder"):
            issues.append(
                {
                    "location": unit.get("unit_id", "unit"),
                    "reason": "存在待核验占位内容，事实性陈述尚未通过 Evidence Gate",
                    "next_step": "补充证据并重新核验后重建该单元",
                }
            )
        for ref in unit.get("claim_refs", []) or []:
            ref_count += 1
            if not ctx.evidence.is_supported(ref):
                try:
                    cid, ver = ref.split("@v", 1)
                    claim = ctx.evidence.get_claim(cid, int(ver))
                    detail = claim.get("status")
                except (KeyError, ValueError, TypeError):
                    detail = "未登记"
                issues.append(
                    {
                        "location": unit.get("unit_id", "unit"),
                        "reason": f"引用了未获支持的 Claim：{ref}（状态 {detail}）",
                        "next_step": "改为待核验占位，或完成证据核验后重建",
                    }
                )
    for artifact in [a for a in ctx.store.list_artifacts() if a["artifact_type"] == "case"]:
        case = ctx.accepted(artifact["artifact_id"]) or {}
        if case.get("section_id") != script.get("section_id"):
            continue
        for ref in case.get("claim_refs", []) or []:
            ref_count += 1
            if not ctx.evidence.is_supported(ref):
                issues.append(
                    {
                        "location": artifact["artifact_id"],
                        "reason": f"案例引用了未获支持的 Claim：{ref}",
                        "next_step": "标注为虚构/教学假设，或完成核验",
                    }
                )
        if case.get("new_factual_claims"):
            issues.append(
                {
                    "location": artifact["artifact_id"],
                    "reason": "案例中存在新增事实性陈述，尚未通过 Evidence Gate",
                    "next_step": "登记为 PENDING Claim 并完成核验",
                }
            )
    status = "FAIL" if issues else "PASS"
    deps: dict[str, str] = {}
    return make_result(
        "G2",
        target_artifact=script_artifact,
        target_rev_id=rev_id,
        target_hash=rev_hash,
        status=status,
        issues=issues,
        dep_versions=deps,
        run_id=ctx.run_id,
        environment=ctx.environment,
    )


def g3_pedagogy(ctx: GateContext, script_artifact: str) -> dict[str, Any]:
    rev_id, rev_hash = _require_rev(ctx, script_artifact)
    script = ctx.accepted(script_artifact) or {}
    section_id = script.get("section_id")
    design = ctx.accepted(f"learning_design:{section_id}") or {}
    plan = ctx.accepted(f"teaching_plan:{section_id}") or {}
    assessment = ctx.accepted(f"assessment:{section_id}") or {}
    issues: list[dict[str, str]] = []
    goals = {g["goal_id"]: g for g in design.get("goals", [])}
    if not goals:
        issues.append(
            {
                "location": f"learning_design:{section_id}",
                "reason": "缺少学习目标",
                "next_step": "补全学习目标后重建",
            }
        )
    if not design.get("assessment_evidence"):
        issues.append(
            {
                "location": f"learning_design:{section_id}:assessment_evidence",
                "reason": "缺少评价证据（目标—评价对应）",
                "next_step": "补充评价证据",
            }
        )
    strategy_by_goal = {s.get("goal_ref"): s for s in plan.get("strategies", [])}
    for goal_id, goal in goals.items():
        strategy = strategy_by_goal.get(goal_id)
        if not strategy:
            issues.append(
                {
                    "location": f"teaching_plan:{section_id}",
                    "reason": f"目标 {goal_id} 没有对应教学策略",
                    "next_step": "补充策略设计",
                }
            )
            continue
        allowed = router.allowed_strategies(goal.get("knowledge_type", "concept"))
        if strategy.get("strategy") not in allowed:
            issues.append(
                {
                    "location": f"teaching_plan:{section_id}:{goal_id}",
                    "reason": f"策略 {strategy.get('strategy')} 不在知识类型 {goal.get('knowledge_type')} 的候选集合 {allowed}",
                    "next_step": "改用候选策略或说明理由",
                }
            )
    if not plan.get("pck_notes"):
        issues.append(
            {
                "location": f"teaching_plan:{section_id}",
                "reason": "缺少 PCK 教学说明",
                "next_step": "补充教学说明",
            }
        )
    for item in assessment.get("items", []):
        if item.get("target_goal") not in goals:
            issues.append(
                {
                    "location": f"assessment:{section_id}:{item.get('assessment_id')}",
                    "reason": f"评价题目指向未知目标：{item.get('target_goal')}",
                    "next_step": "关联到有效学习目标",
                }
            )
    for artifact in [a for a in ctx.store.list_artifacts() if a["artifact_type"] == "case"]:
        case = ctx.accepted(artifact["artifact_id"]) or {}
        if case.get("section_id") == section_id and case.get("linked_goal") not in goals:
            issues.append(
                {
                    "location": artifact["artifact_id"],
                    "reason": f"案例关联的目标不存在：{case.get('linked_goal')}",
                    "next_step": "关联到有效学习目标",
                }
            )
    status = "FAIL" if issues else "PASS"
    deps = {
        k: v
        for k, v in {
            "learning_design": ctx.accepted_rev(f"learning_design:{section_id}"),
            "teaching_plan": ctx.accepted_rev(f"teaching_plan:{section_id}"),
            "assessment": ctx.accepted_rev(f"assessment:{section_id}"),
        }.items()
        if v
    }
    return make_result(
        "G3",
        target_artifact=script_artifact,
        target_rev_id=rev_id,
        target_hash=rev_hash,
        status=status,
        issues=issues,
        dep_versions=deps,
        run_id=ctx.run_id,
        environment=ctx.environment,
    )


def g4_safety(ctx: GateContext, script_artifact: str) -> dict[str, Any]:
    rev_id, rev_hash = _require_rev(ctx, script_artifact)
    script = ctx.accepted(script_artifact) or {}
    section_id = script.get("section_id")
    hard_issues: list[dict[str, str]] = []
    soft_issues: list[dict[str, str]] = []
    segments: list[tuple[str, str]] = list(_units_text(script))
    for artifact in ctx.store.list_artifacts():
        if artifact["artifact_type"] == "case":
            case = ctx.accepted(artifact["artifact_id"]) or {}
            if case.get("section_id") == section_id:
                segments.append((artifact["artifact_id"], case.get("text", "")))
        if artifact["artifact_type"] == "assessment":
            assessment = ctx.accepted(artifact["artifact_id"]) or {}
            if assessment.get("section_id") == section_id:
                for item in assessment.get("items", []):
                    blob = " ".join(
                        [item.get("question", ""), item.get("answer", ""), item.get("rationale", ""), *item.get("options", [])]
                    )
                    segments.append((f"{artifact['artifact_id']}:{item.get('assessment_id')}", blob))
    for pattern in config.RISK.get("safety_patterns", []):
        regex = re.compile(pattern["pattern"])
        for location, text in segments:
            for match in regex.finditer(text or ""):
                hard_issues.append(
                    {
                        "location": location,
                        "reason": f"{pattern['reason']}：命中“{match.group(0)[:30]}”",
                        "next_step": "改写为去诊断、去标签、限定范围的表述",
                    }
                )
    individual = config.RISK.get("individual_inference", {})
    if individual:
        regex = re.compile(individual["pattern"])
        for location, text in segments:
            for match in regex.finditer(text or ""):
                window = (text or "")[max(0, match.start() - 20) : match.end() + 30]
                if any(marker in window for marker in ["教学假设", "可能", "人工复核", "不能据此", "不足以"]):
                    continue
                soft_issues.append(
                    {
                        "location": location,
                        "reason": f"{individual['reason']}：命中“{match.group(0)}”",
                        "next_step": "标注为教学假设，或改为群体层面表述",
                    }
                )
    if hard_issues:
        status = "FAIL"
    elif soft_issues:
        status = "NEEDS_REVIEW"
    else:
        status = "PASS"
    return make_result(
        "G4",
        target_artifact=script_artifact,
        target_rev_id=rev_id,
        target_hash=rev_hash,
        status=status,
        issues=hard_issues + soft_issues,
        run_id=ctx.run_id,
        environment=ctx.environment,
    )


def _udl_strings(value: Any, path: str = "udl") -> list[tuple[str, str]]:
    """Every string inside a UDL payload, with a location for the issue.

    The payload nests barriers, options and review flags, and the owner spec
    names some of those fields differently from the schema, so the ban is applied
    to the text wherever it sits rather than to a guessed field name.
    """
    found: list[tuple[str, str]] = []
    if isinstance(value, str):
        found.append((path, value))
    elif isinstance(value, dict):
        for key, item in value.items():
            found.extend(_udl_strings(item, f"{path}.{key}"))
    elif isinstance(value, (list, tuple)):
        for index, item in enumerate(value):
            found.extend(_udl_strings(item, f"{path}[{index}]"))
    return found


def g5_udl(ctx: GateContext, script_artifact: str) -> dict[str, Any]:
    rev_id, rev_hash = _require_rev(ctx, script_artifact)
    script = ctx.accepted(script_artifact) or {}
    section_id = script.get("section_id")
    design = ctx.accepted(f"learning_design:{section_id}") or {}
    plan = ctx.accepted(f"teaching_plan:{section_id}") or {}
    goals = {g["goal_id"] for g in design.get("goals", [])}
    udl = plan.get("udl") or {}
    issues: list[dict[str, str]] = []
    if not udl:
        issues.append(
            {
                "location": f"teaching_plan:{section_id}:udl",
                "reason": "缺少 UDL 障碍与替代路径分析",
                "next_step": "运行 UDL 检查步骤",
            }
        )
    else:
        if not udl.get("barriers"):
            issues.append(
                {
                    "location": f"teaching_plan:{section_id}:udl.barriers",
                    "reason": "未列出学习障碍",
                    "next_step": "补充障碍分析",
                }
            )
        if not udl.get("options"):
            issues.append(
                {
                    "location": f"teaching_plan:{section_id}:udl.options",
                    "reason": "未提供替代路径",
                    "next_step": "补充替代路径",
                }
            )
        for option in udl.get("options", []):
            if not option.get("core_demand_preserved", False):
                issues.append(
                    {
                        "location": f"teaching_plan:{section_id}:udl",
                        "reason": "替代路径降低了核心认知要求（将更简单任务伪装成更无障碍任务）",
                        "next_step": "保持同一目标与核心认知要求",
                    }
                )
            if option.get("same_goal_ref") not in goals:
                issues.append(
                    {
                        "location": f"teaching_plan:{section_id}:udl",
                        "reason": f"替代路径未关联有效学习目标：{option.get('same_goal_ref')}",
                        "next_step": "关联到有效目标",
                    }
                )
            if not any(option.get(k) for k in ["representation", "action_expression", "engagement"]):
                issues.append(
                    {
                        "location": f"teaching_plan:{section_id}:udl",
                        "reason": "替代路径未说明 UDL 维度（表征/行动表达/参与）",
                        "next_step": "补充 UDL 维度说明",
                    }
                )
        # Section 27 bans referral wording outright. It is a deterministic policy
        # violation, so it fails the gate: degrading it to a warning or review
        # state would let the same wording reach the learner anyway.
        for location, text in _udl_strings(udl):
            if any(re.search(pattern, text) for pattern in policies.REFERRAL_PATTERNS):
                issues.append(
                    {
                        "location": f"teaching_plan:{section_id}:{location}",
                        "reason": f"出现禁止的转介措辞：{text[:40]}",
                        "next_step": "改为课堂内可执行的支持，不依赖转介",
                    }
                )
    status = "FAIL" if issues else "PASS"
    deps = {k: v for k, v in {"teaching_plan": ctx.accepted_rev(f"teaching_plan:{section_id}")}.items() if v}
    return make_result(
        "G5",
        target_artifact=script_artifact,
        target_rev_id=rev_id,
        target_hash=rev_hash,
        status=status,
        issues=issues,
        dep_versions=deps,
        run_id=ctx.run_id,
        environment=ctx.environment,
    )


def g6_multimedia(ctx: GateContext, script_artifact: str) -> dict[str, Any]:
    rev_id, rev_hash = _require_rev(ctx, script_artifact)
    script = ctx.accepted(script_artifact) or {}
    section_id = script.get("section_id")
    reqs = ctx.accepted("requirements") or {}
    section = next((s for s in reqs.get("sections", []) if s.get("section_id") == section_id), {})
    media_plan = ctx.accepted(f"media_plan:{section_id}")
    diagrams = ctx.accepted(f"diagrams:{section_id}") or {"diagrams": []}
    decisions_doc = ctx.accepted(f"animation_decisions:{section_id}") or {"decisions": []}
    storyboard = ctx.accepted(f"storyboard:{section_id}") or {}
    decisions = decisions_doc.get("decisions", [])
    approved = [d for d in decisions if d.get("decision") == "ANIMATION"]
    animation_required = "animation" in section.get("required_components", [])
    issues: list[dict[str, str]] = []
    needs_review: list[dict[str, str]] = []

    for unit in script.get("units", []):
        if unit.get("kind") == "visual" and not unit.get("visual_function"):
            issues.append(
                {
                    "location": unit.get("unit_id", "visual"),
                    "reason": "画面提示未说明知识功能（视觉是否承担知识功能）",
                    "next_step": "补充 visual_function",
                }
            )
    if media_plan is None:
        needs_review.append(
            {
                "location": f"media_plan:{section_id}",
                "reason": "媒体选型未运行，无法确认视觉与动画决策",
                "next_step": "运行媒体选型与动画 Gate 后重跑",
            }
        )
    else:
        for item in media_plan.get("items", []):
            if item.get("recommended_medium") == "diagram" and not any(
                diagram.get("item_id") == item.get("item_id") for diagram in diagrams.get("diagrams", [])
            ):
                issues.append(
                    {
                        "location": f"diagrams:{section_id}:{item.get('item_id')}",
                        "reason": f"媒体选型推荐图示但未生成对应 SVG 资产（{item.get('item_id')}）",
                        "next_step": "重新生成图示",
                    }
                )
    load_review = ctx.accepted(f"load_review:{section_id}") or {}
    cognitive_load = load_review.get("cognitive_load") or {}
    if cognitive_load.get("level") == "HIGH":
        suggestions = "；".join(str(item) for item in cognitive_load.get("improvements", [])[:3])
        needs_review.append(
            {
                "location": f"load_review:{section_id}",
                "reason": "认知负荷评估为 HIGH：" + "；".join(str(item) for item in cognitive_load.get("reasons", [])[:3]),
                "next_step": suggestions or "分段、去掉冗余、为视觉加信号标注后重跑",
            }
        )
    for pairing in (load_review.get("dual_coding") or {}).get("pairings", []):
        if pairing.get("pairing_ok") is False:
            issues.append(
                {
                    "location": f"load_review:{section_id}:{pairing.get('goal_ref')}",
                    "reason": "双重编码配对未通过：视觉没有承担知识功能",
                    "next_step": "让视觉承担 structure/comparison 等明确功能，或改为文本",
                }
            )
    if animation_required:
        if not decisions:
            needs_review.append(
                {
                    "location": f"requirements:{section_id}",
                    "reason": "要求包含动画，但 Animation Gate 没有任何候选或决策",
                    "next_step": "补充可动画的知识点，或修订动画要求（第 58 章冲突处理）",
                }
            )
        elif not approved:
            for decision in decisions:
                needs_review.append(
                    {
                        "location": f"animation:{decision.get('item_id')}",
                        "reason": "未通过 Animation Gate：" + str(decision.get("rationale", "")),
                        "next_step": "静态替代：" + str(decision.get("static_alternative_desc", "见媒体计划")) + "；确认后可修订动画要求",
                    }
                )
    if approved:
        if not storyboard.get("shots"):
            issues.append(
                {
                    "location": f"storyboard:{section_id}",
                    "reason": "存在已批准的动画，但缺少教学分镜与生成提示词",
                    "next_step": "重新运行动画决策与分镜步骤",
                }
            )
        else:
            if not str(storyboard.get("static_terminal_state", "")).strip():
                issues.append(
                    {
                        "location": f"storyboard:{section_id}",
                        "reason": "动画缺少静态终态设计（学生需要反复查看的终态）",
                        "next_step": "补充 static_terminal_state",
                    }
                )
            if storyboard.get("rendered"):
                issues.append(
                    {
                        "location": f"storyboard:{section_id}",
                        "reason": "分镜标记为已渲染；本系统必须在视频生成前停止",
                        "next_step": "移除渲染标记，交由外部视频系统执行",
                    }
                )
            if any(not str(shot.get("narration", "")).strip() for shot in storyboard.get("shots", [])):
                issues.append(
                    {
                        "location": f"storyboard:{section_id}",
                        "reason": "存在缺少讲解同步文本的镜头",
                        "next_step": "补充每个镜头的 narration",
                    }
                )
    if issues:
        status = "FAIL"
    elif needs_review:
        status = "NEEDS_REVIEW"
    elif approved or animation_required:
        status = "PASS"
    else:
        status = "NOT_APPLICABLE"
    return make_result(
        "G6",
        target_artifact=script_artifact,
        target_rev_id=rev_id,
        target_hash=rev_hash,
        status=status,
        issues=issues + needs_review,
        run_id=ctx.run_id,
        environment=ctx.environment,
        na_reason=None
        if status != "NOT_APPLICABLE"
        else "无动画要求且无批准的动画候选；媒体为文本/静态图示",
    )


def g7_template(ctx: GateContext, script_artifact: str) -> dict[str, Any]:
    rev_id, rev_hash = _require_rev(ctx, script_artifact)
    script = ctx.accepted(script_artifact) or {}
    spec = ctx.accepted("template_spec") or {}
    if spec.get("kind") in (None, "none") or not spec.get("columns"):
        return make_result(
            "G7",
            target_artifact=script_artifact,
            target_rev_id=rev_id,
            target_hash=rev_hash,
            status="NOT_APPLICABLE",
            run_id=ctx.run_id,
            environment=ctx.environment,
            na_reason="未提供可解析的模板结构",
        )
    issues: list[dict[str, str]] = []
    expected_ref = content_hash(spec)
    if script.get("template_ref") != expected_ref:
        issues.append(
            {
                "location": f"{script_artifact}:template_ref",
                "reason": "脚本未按当前模板版本生成（template_ref 不匹配）",
                "next_step": "按当前模板重建脚本",
            }
        )
    text = _all_script_text(script)
    for rule in spec.get("terminology_rules", []):
        if rule and rule not in text:
            issues.append(
                {
                    "location": f"{script_artifact}:terminology",
                    "reason": f"模板术语规则未满足：未出现“{rule}”",
                    "next_step": f"按模板使用“{rule}”",
                }
            )
    for column in spec.get("columns", []):
        name = column.get("name", "")
        kind = next((k for k, words in COLUMN_KIND_MAP.items() if any(w in name for w in words)), None)
        if kind and not _script_units_of_kind(script, kind):
            issues.append(
                {
                    "location": f"{script_artifact}:columns",
                    "reason": f"模板栏目“{name}”没有对应的脚本单元（{kind}）",
                    "next_step": "按模板栏目补充内容",
                }
            )
    status = "FAIL" if issues else "PASS"
    deps = {k: v for k, v in {"template_spec": ctx.accepted_rev("template_spec")}.items() if v}
    return make_result(
        "G7",
        target_artifact=script_artifact,
        target_rev_id=rev_id,
        target_hash=rev_hash,
        status=status,
        issues=issues,
        dep_versions=deps,
        run_id=ctx.run_id,
        environment=ctx.environment,
    )


def g8_artifact(ctx: GateContext, export_artifact: str) -> dict[str, Any]:
    rev_id, rev_hash = _require_rev(ctx, export_artifact)
    manifest = ctx.accepted(export_artifact) or {}
    issues: list[dict[str, str]] = []
    soft_issues: list[dict[str, str]] = []
    from pathlib import Path

    from . import media as media_mod

    for file_entry in manifest.get("files", []):
        path = Path(file_entry["path"])
        role = file_entry.get("role") or "lesson_script"
        kind = file_entry.get("kind")
        if not path.exists():
            issues.append(
                {
                    "location": file_entry["path"],
                    "reason": "导出文件不存在",
                    "next_step": "重新导出",
                }
            )
            continue
        if kind == "docx":
            try:
                import docx

                document = docx.Document(str(path))
                paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
                if not paragraphs and not document.tables:
                    issues.append(
                        {
                            "location": file_entry["path"],
                            "reason": "DOCX 可打开但没有可读内容",
                            "next_step": "检查导出模板与内容",
                        }
                    )
                if role in LABEL_CHECK_ROLES and manifest.get("mode") == "formal" and document.tables:
                    for ti, table in enumerate(document.tables):
                        for ri, row in enumerate(table.rows):
                            cells = [c.text for c in row.cells]
                            if any("草稿" in c for c in cells):
                                issues.append(
                                    {
                                        "location": f"{file_entry['path']}:table{ti}:row{ri}",
                                        "reason": "正式文件中出现草稿标识",
                                        "next_step": "重新生成正式导出",
                                    }
                                )
            except Exception as exc:  # noqa: BLE001 - open failure must be reported, not swallowed
                issues.append(
                    {
                        "location": file_entry["path"],
                        "reason": f"DOCX 无法打开或解析：{exc}",
                        "next_step": "检查导出实现与文件完整性",
                    }
                )
        elif kind == "markdown":
            content = path.read_text(encoding="utf-8")
            if role in LABEL_CHECK_ROLES:
                if manifest.get("mode") == "draft" and "草稿—未通过 QA" not in content:
                    issues.append(
                        {
                            "location": file_entry["path"],
                            "reason": "草稿 Markdown 缺少草稿标识",
                            "next_step": "在导出中加入草稿标注",
                        }
                    )
                if manifest.get("mode") == "formal" and "草稿—未通过 QA" in content:
                    issues.append(
                        {
                            "location": file_entry["path"],
                            "reason": "正式文件含草稿标识",
                            "next_step": "重新生成正式导出",
                        }
                    )
            if role == "worksheet" and "答案：" in content:
                issues.append(
                    {
                        "location": file_entry["path"],
                        "reason": "学生版学习单包含教师答案",
                        "next_step": "重新导出学生版（不含答案）",
                    }
                )
        elif kind == "svg":
            content = path.read_text(encoding="utf-8")
            for problem in media_mod.validate_svg(content):
                issues.append(
                    {
                        "location": file_entry["path"],
                        "reason": f"SVG 校验未通过：{problem}",
                        "next_step": "重新生成图示",
                    }
                )
        elif kind == "json":
            import json as _json

            try:
                data = _json.loads(path.read_text(encoding="utf-8"))
            except ValueError as exc:
                issues.append(
                    {
                        "location": file_entry["path"],
                        "reason": f"JSON 解析失败：{exc}",
                        "next_step": "重新生成该文件",
                    }
                )
                data = None
            if isinstance(data, dict) and role == "storyboard" and data.get("rendered"):
                issues.append(
                    {
                        "location": file_entry["path"],
                        "reason": "分镜文件标记为已渲染；本系统必须在视频生成前停止",
                        "next_step": "移除渲染标记",
                    }
                )
        elif kind == "txt":
            content = path.read_text(encoding="utf-8")
            if not content.strip():
                issues.append(
                    {
                        "location": file_entry["path"],
                        "reason": "视频生成提示词为空",
                        "next_step": "重新生成提示词",
                    }
                )
        elif kind == "pptx":
            try:
                from pptx import Presentation

                deck = Presentation(str(path))
                if len(deck.slides) == 0:
                    issues.append(
                        {
                            "location": file_entry["path"],
                            "reason": "PPTX 可打开但没有幻灯片",
                            "next_step": "检查幻灯片沟通计划与生成步骤",
                        }
                    )
            except Exception as exc:  # noqa: BLE001 - open failure must be reported, not swallowed
                issues.append(
                    {
                        "location": file_entry["path"],
                        "reason": f"PPTX 无法打开或解析：{exc}",
                        "next_step": "检查 PPTX 生成实现",
                    }
                )
            deck_artifact = ctx.accepted("pptx_deck") or {}
            from . import render as render_mod

            renderer = render_mod.find_renderer()
            if not renderer:
                soft_issues.append(
                    {
                        "location": file_entry["path"],
                        "reason": "无可用渲染器（PowerPoint/LibreOffice），无法完成渲染级 QA（83.4）",
                        "next_step": "安装渲染器后重跑导出",
                    }
                )
            elif not deck_artifact.get("thumbnails"):
                soft_issues.append(
                    {
                        "location": file_entry["path"],
                        "reason": "渲染器可用但未生成缩略图，渲染级检查未完成",
                        "next_step": "重跑 PPT 生成步骤",
                    }
                )
            qa = deck_artifact.get("qa") or {}
            if qa.get("status") == "FAIL":
                for qa_issue in qa.get("issues", []):
                    issues.append(
                        {
                            "location": qa_issue.get("location", str(path)),
                            "reason": qa_issue.get("reason", "PPTX 检查失败"),
                            "next_step": qa_issue.get("next_step", "修正后重跑"),
                        }
                    )
            elif qa.get("status") == "NEEDS_REVIEW":
                reasons = "；".join(str(item.get("reason", "")) for item in qa.get("issues", [])[:3])
                soft_issues.append(
                    {
                        "location": file_entry["path"],
                        "reason": "PPTX 结构检查完成但存在需复核项：" + (reasons or "渲染级检查未完成"),
                        "next_step": "安装 LibreOffice/PowerPoint 或人工复核；在此之前正式导出保持阻断",
                    }
                )
    status = "FAIL" if issues else ("NEEDS_REVIEW" if soft_issues else "PASS")
    try:
        from importlib.metadata import version as _pkg_version

        renderer = f"python-docx {_pkg_version('python-docx')} · 字体 Microsoft YaHei"
    except Exception:  # noqa: BLE001 - renderer metadata is informational
        renderer = "python-docx · 字体 Microsoft YaHei"
    return make_result(
        "G8",
        target_artifact=export_artifact,
        target_rev_id=rev_id,
        target_hash=rev_hash,
        status=status,
        issues=issues + soft_issues,
        dep_versions={},
        run_id=ctx.run_id,
        environment=ctx.environment,
        renderer=renderer,
    )


GATE_FUNCTIONS: dict[str, Callable[[GateContext, str], dict[str, Any]]] = {
    "G1": g1_requirements,
    "G2": g2_evidence,
    "G3": g3_pedagogy,
    "G4": g4_safety,
    "G5": g5_udl,
    "G6": g6_multimedia,
    "G7": g7_template,
    "G8": g8_artifact,
}


def run_section_gates(ctx: GateContext, script_artifact: str) -> list[dict[str, Any]]:
    return [GATE_FUNCTIONS[gate_id](ctx, script_artifact) for gate_id in GATE_ORDER]


def write_gate_results(store: Store, results: list[dict[str, Any]], run_id: str) -> list[str]:
    written: list[str] = []
    for result in results:
        artifact_id = gate_artifact_id(result["gate_id"], result["target"]["artifact_id"])
        info = store.add_revision(
            artifact_id=artifact_id,
            artifact_type="gate_result",
            content=result,
            produced_by="gate-runner",
            run_id=run_id,
            rules_version=config.RULES_VERSION,
            deps=[result["target"]["revision_id"], *result.get("dep_versions", {}).values()],
        )
        written.append(info["revision_id"])
    return written


def latest_gate_results(store: Store, overrides: dict[str, str] | None = None) -> dict[str, dict[str, Any]]:
    overrides = overrides or {}
    results: dict[str, dict[str, Any]] = {}
    for artifact in store.list_artifacts():
        if artifact["artifact_type"] != "gate_result":
            continue
        artifact_id = artifact["artifact_id"]
        rev_id = overrides.get(artifact_id) or store.accepted_rev_id(artifact_id)
        if rev_id:
            results[artifact_id] = store.get_revision(rev_id)["content"]
    return results


def export_readiness(store: Store, overrides: dict[str, str] | None = None) -> dict[str, Any]:
    artifacts = store.list_artifacts()
    scripts = [a for a in artifacts if a["artifact_type"] == "script"]
    results = latest_gate_results(store, overrides)
    blocking: list[str] = []
    summary: list[dict[str, Any]] = []
    if not scripts:
        blocking.append("尚无脚本产物")
    for script in scripts:
        artifact_id = script["artifact_id"]
        rev_id = (overrides or {}).get(artifact_id) or script["accepted_rev"]
        if script["stale"] and artifact_id not in (overrides or {}):
            blocking.append(f"{artifact_id} 已过期，需重建后再导出")
        if not rev_id:
            blocking.append(f"{artifact_id} 没有已接受的版本")
            continue
        try:
            rev_row = store.get_revision(rev_id)
        except StoreError:
            rev_row = None
        if rev_row and rev_row.get("fixture"):
            blocking.append(f"{artifact_id} 的最新内容来自测试夹具（test_fixture），不能正式导出")
        for gate_id in GATE_ORDER:
            key = gate_artifact_id(gate_id, artifact_id)
            result = results.get(key)
            entry = {"gate_id": gate_id, "artifact_id": artifact_id, "status": None, "current": False}
            if result is None:
                blocking.append(f"{gate_id} 尚未在 {artifact_id} 上运行")
                summary.append(entry)
                continue
            current = (
                result["target"]["revision_id"] == rev_id
                and result["rules_version"] == config.RULES_VERSION
            )
            entry["status"] = result["status"]
            entry["current"] = current
            if not current:
                blocking.append(f"{gate_id} 结果已过期（目标版本或规则版本变化）")
            elif result["status"] not in PASSING:
                blocking.append(f"{gate_id} 状态为 {result['status']}")
            elif result["status"] == "NOT_APPLICABLE" and not result.get("not_applicable_reason"):
                blocking.append(f"{gate_id} 的 NOT_APPLICABLE 缺少原因")
            summary.append(entry)
    return {"ready": not blocking, "blocking": blocking, "summary": summary}


def export_readiness_with_overrides(ctx: GateContext) -> dict[str, Any]:
    return export_readiness(ctx.store, ctx.overrides)
