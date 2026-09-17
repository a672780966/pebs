from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
TEXT_DARK = RGBColor(0x1F, 0x24, 0x30)
TEXT_MUTED = RGBColor(0x6B, 0x72, 0x80)
DEFAULT_ACCENT = RGBColor(0x2B, 0x3D, 0x7A)
BOX_FILL = RGBColor(0xEE, 0xF1, 0xF8)
BOX_FILL_ALT = RGBColor(0xE8, 0xF4, 0xEE)
FONT = "Microsoft YaHei"

INTERNAL_JARGON = [
    "claim_refs",
    "claim_id",
    "artifact",
    "Animation Gate",
    "Media Router",
    "step_id",
    "provider",
    "run_",
    "clm_",
    "cs_",
    "QA",
    "DAG",
    "diff",
]

_FALLBACK_FAMILY = {"title_size": 28, "body_size": 18, "accent": "2B3D7A"}


def _families() -> dict[str, Any]:
    return getattr(__import__("pebs.config", fromlist=["SLIDE_FAMILIES"]), "SLIDE_FAMILIES", {}) or {}


def _style_for(row: dict[str, Any]) -> dict[str, Any]:
    data = _families()
    mapping = data.get("archetype_defaults", {})
    family = str(row.get("family") or mapping.get(row.get("layout_archetype", ""), "title_content"))
    style = dict(_FALLBACK_FAMILY)
    style.update(data.get("slide_families", {}).get(family, {}))
    style["family"] = family
    return style


def _accent(style: dict[str, Any]) -> RGBColor:
    try:
        return RGBColor.from_string(str(style.get("accent", "2B3D7A")))
    except ValueError:
        return DEFAULT_ACCENT


def _set_font(run, size: float, *, bold: bool = False, color: RGBColor = TEXT_DARK) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = rpr.makeelement(qn("a:ea"), {})
        rpr.append(ea)
    ea.set("typeface", FONT)


def _textbox(slide, left: Emu, top: Emu, width: Emu, height: Emu, text: str, size: float, *, bold: bool = False, color: RGBColor = TEXT_DARK, align=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    _set_font(run, size, bold=bold, color=color)
    return box


def _bullets(slide, left: Emu, top: Emu, width: Emu, height: Emu, points: list[str], size: float = 20, *, color: RGBColor = TEXT_DARK):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    for index, point in enumerate(points):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(10)
        run = paragraph.add_run()
        run.text = f"• {point}"
        _set_font(run, size, color=color)
    return box


def _fill_placeholder(slide, kind: str, text: str, size: float, *, bold: bool = False, bullets: bool = False) -> bool:
    for placeholder in slide.placeholders:
        if kind == "title" and placeholder.placeholder_format.idx != 0:
            continue
        if kind == "body" and placeholder.placeholder_format.idx in (0,):
            continue
        frame = placeholder.text_frame
        frame.word_wrap = True
        if bullets:
            lines = text.splitlines() or [""]
            for index, line in enumerate(lines):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                run = paragraph.add_run()
                run.text = f"• {line}"
                _set_font(run, size)
        else:
            frame.text = text
            for paragraph in frame.paragraphs:
                for run in paragraph.runs:
                    _set_font(run, size, bold=bold)
        return True
    return False


def _add_page_number(slide, number: int, citations: list[str]) -> None:
    footer = f"{number}"
    if citations:
        footer = f"{number} · 引用：{'、'.join(citations[:4])}"
    _textbox(slide, Inches(0.6), SLIDE_H - Inches(0.55), SLIDE_W - Inches(1.2), Inches(0.4), footer, 10, color=TEXT_MUTED)


def _diagram_shapes(slide, semantics: dict[str, Any], top: Emu, accent: RGBColor) -> None:
    nodes = semantics.get("nodes", [])[:8]
    edges = semantics.get("edges", [])
    if not nodes:
        return
    columns = 2
    box_w = Inches(5.2)
    box_h = Inches(1.1)
    gap_x = Inches(0.7)
    gap_y = Inches(0.6)
    start_x = Inches(0.8)
    positions: dict[str, tuple[Emu, Emu]] = {}
    for index, node in enumerate(nodes):
        row, column = divmod(index, columns)
        left = start_x + column * (box_w + gap_x)
        box_top = top + row * (box_h + gap_y)
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, box_top, box_w, box_h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = BOX_FILL if column == 0 else BOX_FILL_ALT
        shape.line.color.rgb = accent
        frame = shape.text_frame
        frame.word_wrap = True
        paragraph = frame.paragraphs[0]
        run = paragraph.add_run()
        run.text = str(node.get("label", node.get("id", "")))
        _set_font(run, 16, bold=True)
        positions[str(node.get("id"))] = (left + box_w // 2, box_top + box_h // 2)
    for edge in edges[:12]:
        start = positions.get(str(edge.get("from")))
        end = positions.get(str(edge.get("to")))
        if not start or not end:
            continue
        connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, start[0], start[1], end[0], end[1])
        connector.line.color.rgb = accent
        if edge.get("label"):
            mid_left = min(start[0], end[0]) + abs(end[0] - start[0]) // 2
            mid_top = min(start[1], end[1]) + abs(end[1] - start[1]) // 2
            _textbox(slide, mid_left - Inches(0.6), mid_top - Inches(0.18), Inches(1.6), Inches(0.36), str(edge["label"]), 11, color=TEXT_MUTED)


def _pick_layout(presentation: Presentation, keywords: list[str], fallback_index: int):
    for layout in presentation.slide_layouts:
        name = str(getattr(layout, "name", ""))
        if any(keyword.lower() in name.lower() for keyword in keywords):
            return layout
    layouts = list(presentation.slide_layouts)
    if fallback_index < len(layouts):
        return layouts[fallback_index]
    return layouts[0] if layouts else None


def build_deck(
    *,
    course: str,
    plan_rows: list[dict[str, Any]],
    output_path: Path,
    template_path: Path | None = None,
    use_template: bool = True,
) -> dict[str, Any]:
    template_used = ""
    if template_path and use_template and Path(template_path).suffix.lower() == ".pptx" and Path(template_path).exists():
        presentation = Presentation(str(template_path))
        template_used = str(template_path)
    else:
        presentation = Presentation()
        presentation.slide_width = SLIDE_W
        presentation.slide_height = SLIDE_H
    blank = _pick_layout(presentation, ["空白", "Blank"], 6)
    title_layout = _pick_layout(presentation, ["标题幻灯片", "Title Slide", "标题"], 0)
    content_layout = _pick_layout(presentation, ["标题和内容", "Title and Content", "内容"], 1)

    stats = {"slides": 0, "text_boxes": 0, "shapes": 0, "tables": 0, "connectors": 0, "pictures": 0}
    slide_records: list[dict[str, Any]] = []

    for index, row in enumerate(plan_rows, start=1):
        archetype = row.get("layout_archetype", "bullets")
        style = _style_for(row)
        accent = _accent(style)
        title = str(row.get("title", ""))
        citations = [str(ref) for ref in row.get("citation_labels", [])]
        layout = title_layout if archetype in ("cover", "section_title", "closing") else content_layout
        slide = presentation.slides.add_slide(layout) if layout is not None else presentation.slides.add_slide(blank)
        if archetype == "cover":
            if not _fill_placeholder(slide, "title", title, style["title_size"], bold=True):
                _textbox(slide, Inches(1.0), Inches(2.3), SLIDE_W - Inches(2.0), Inches(1.4), title, style["title_size"], bold=True)
            subtitle = course
            if not _fill_placeholder(slide, "body", subtitle, style["body_size"]):
                _textbox(slide, Inches(1.0), Inches(3.8), SLIDE_W - Inches(2.0), Inches(0.8), subtitle, style["body_size"], color=TEXT_MUTED)
        elif archetype == "section_title":
            if not _fill_placeholder(slide, "title", title, style["title_size"], bold=True):
                _textbox(slide, Inches(1.0), Inches(2.8), SLIDE_W - Inches(2.0), Inches(1.2), title, style["title_size"], bold=True)
            if row.get("core_message"):
                _textbox(slide, Inches(1.0), Inches(4.1), SLIDE_W - Inches(2.0), Inches(0.8), str(row["core_message"]), style["body_size"], color=TEXT_MUTED)
        else:
            if not _fill_placeholder(slide, "title", title, style["title_size"], bold=True):
                _textbox(slide, Inches(0.8), Inches(0.5), SLIDE_W - Inches(1.6), Inches(0.9), title, style["title_size"], bold=True)
            if archetype == "diagram":
                _diagram_shapes(slide, row.get("diagram_semantics") or {}, Inches(1.8), accent)
            elif archetype == "table" and row.get("table"):
                headers = row["table"].get("headers", [])
                rows = row["table"].get("rows", [])
                table_shape = slide.shapes.add_table(len(rows) + 1, max(1, len(headers)), Inches(0.8), Inches(1.8), SLIDE_W - Inches(1.6), Inches(0.5) * (len(rows) + 1))
                table = table_shape.table
                stats["tables"] += 1
                for column, header in enumerate(headers):
                    cell = table.cell(0, column)
                    cell.text = str(header)
                    for paragraph in cell.text_frame.paragraphs:
                        for run in paragraph.runs:
                            _set_font(run, style["body_size"] + 1, bold=True)
                for row_index, data_row in enumerate(rows, start=1):
                    for column, value in enumerate(data_row[: len(headers)]):
                        cell = table.cell(row_index, column)
                        cell.text = str(value)
                        for paragraph in cell.text_frame.paragraphs:
                            for run in paragraph.runs:
                                _set_font(run, style["body_size"])
            else:
                points = [str(point) for point in row.get("points", [])][:6]
                _bullets(slide, Inches(0.9), Inches(1.7), SLIDE_W - Inches(1.8), SLIDE_H - Inches(2.6), points, style["body_size"])
        notes = str(row.get("notes", "")).strip()
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        _add_page_number(slide, index, citations)
        stats["slides"] += 1
        stats["shapes"] += sum(1 for shape in slide.shapes if shape.shape_type is not None)
        stats["text_boxes"] += sum(1 for shape in slide.shapes if shape.has_text_frame)
        stats["connectors"] += sum(1 for shape in slide.shapes if str(shape.element.tag).endswith("cxnSp"))
        stats["pictures"] += sum(1 for shape in slide.shapes if str(shape.element.tag).endswith("pic"))
        slide_records.append(
            {
                "slide_id": f"slide{index}",
                "title": title,
                "layout": archetype,
                "family": style["family"],
                "section_id": row.get("section_id", ""),
                "citation_labels": citations,
                "notes_chars": len(notes),
            }
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(output_path))
    return {
        "path": str(output_path),
        "stats": stats,
        "slides": slide_records,
        "template_used": template_used,
        "slide_width_in": presentation.slide_width / 914400,
        "slide_height_in": presentation.slide_height / 914400,
    }


def inspect_deck(path: Path) -> dict[str, Any]:
    presentation = Presentation(str(path))
    jargon_hits: list[dict[str, str]] = []
    overflow_suspects: list[dict[str, str]] = []
    slides: list[dict[str, Any]] = []
    for index, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text
            texts.append(text)
            for token in INTERNAL_JARGON:
                if token in text:
                    jargon_hits.append({"location": f"slide{index}", "token": token})
            area = (shape.width / 914400) * (shape.height / 914400)
            chars = len(text.strip())
            if area > 0 and chars > area * 95:
                overflow_suspects.append(
                    {"location": f"slide{index}", "chars": str(chars), "area_in2": f"{area:.1f}"}
                )
        slides.append({"slide_id": f"slide{index}", "texts": texts, "shape_count": len(slide.shapes)})
    return {
        "slide_count": len(slides),
        "slides": slides,
        "jargon_hits": jargon_hits,
        "overflow_suspects": overflow_suspects,
    }
