from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any

from . import gates
from .gates import COLUMN_KIND_MAP
from .store import current_signoffs, file_hash, now_iso

LABEL_FORMAL = "正式"
LABEL_DRAFT = "草稿—未通过 QA"


def _set_cjk_font(document: Any) -> None:
    from docx.oxml.ns import qn

    style = document.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def _column_text(units: list[dict[str, Any]], column_name: str) -> str:
    for kind, words in COLUMN_KIND_MAP.items():
        if any(word in column_name for word in words):
            matched = [u for u in units if u.get("kind") == kind]
            if kind == "visual":
                return "\n".join(f"{u.get('text', '')}（{u.get('visual_function', '')}）" for u in matched)
            return "\n".join(u.get("text", "") for u in matched)
    return ""


def _units_markdown(script: dict[str, Any]) -> str:
    lines: list[str] = []
    for unit in script.get("units", []):
        kind = unit.get("kind")
        text = unit.get("text", "")
        if kind == "title":
            lines.append(f"### {text}")
        elif kind == "visual":
            lines.append(f"> [画面] {text}（{unit.get('visual_function', '')}）")
        elif kind == "note":
            lines.append(f"> {text}")
        elif kind == "case":
            lines.append(f"**案例**：{text}")
        elif kind == "interaction":
            lines.append(f"**互动**：{text}")
        else:
            lines.append(text)
        lines.append("")
    return "\n".join(lines)


def _section_markdown(course: str, script: dict[str, Any], label: str | None, blocking: list[str]) -> str:
    parts = [f"# {course} — {script.get('title', '')}"]
    if label:
        parts.append(f"> **{label}**")
    parts.append("")
    wc = script.get("word_count", {})
    if wc:
        parts.append(
            f"*正文字数：{wc.get('count', 0)}（要求 {wc.get('min', 0)}–{wc.get('max', 0)}；口径：{wc.get('scope_note', '')}）*"
        )
        parts.append("")
    parts.append(_units_markdown(script))
    if blocking:
        parts.append("## 未通过 QA 的问题（草稿）")
        for item in blocking:
            parts.append(f"- {item}")
    return "\n".join(parts)


def _write_docx(path: Path, course: str, script: dict[str, Any], columns: list[str], label: str | None, blocking: list[str]) -> None:
    import docx

    document = docx.Document()
    _set_cjk_font(document)
    if label:
        document.add_paragraph(label)
    document.add_heading(f"{course} — {script.get('title', '')}", level=1)
    wc = script.get("word_count", {})
    if wc:
        document.add_paragraph(
            f"正文字数：{wc.get('count', 0)}（要求 {wc.get('min', 0)}–{wc.get('max', 0)}）"
        )
    units = script.get("units", [])
    if columns:
        table = document.add_table(rows=1, cols=len(columns))
        table.style = "Table Grid"
        for idx, name in enumerate(columns):
            table.rows[0].cells[idx].text = name
        row = table.add_row()
        for idx, name in enumerate(columns):
            row.cells[idx].text = _column_text(units, name)
    else:
        for unit in units:
            kind = unit.get("kind")
            text = unit.get("text", "")
            if kind == "title":
                document.add_heading(text, level=2)
            elif kind == "visual":
                paragraph = document.add_paragraph()
                run = paragraph.add_run(f"【画面】{text}（{unit.get('visual_function', '')}）")
                run.italic = True
            elif kind == "case":
                document.add_paragraph(f"【案例】{text}")
            elif kind == "interaction":
                document.add_paragraph(f"【互动】{text}")
            elif kind == "note":
                document.add_paragraph(text)
            else:
                document.add_paragraph(text)
    if blocking:
        document.add_heading("未通过 QA 的问题（草稿）", level=2)
        for item in blocking:
            document.add_paragraph(item, style="List Bullet")
    document.save(str(path))


def _output_dir(ctx: Any, mode: str) -> Path:
    base = Path(ctx.store.base_dir) / "outputs"
    if mode == "formal":
        return base / "staging" / str(ctx.changeset_id)
    return base / "draft"


def _lesson_markdown(course: str, lesson: dict[str, Any], label: str | None) -> str:
    lines = [f"# {course} — {lesson.get('title', '')}（教案）"]
    if label:
        lines.append(f"> **{label}**")
    lines.append("")
    lines.append("## 教学目标")
    for objective in lesson.get("objectives", []):
        lines.append(f"- （{objective.get('goal_ref', '')}）{objective.get('text', '')}")
    if lesson.get("key_points"):
        lines.append("")
        lines.append("## 教学重点")
        for item in lesson["key_points"]:
            lines.append(f"- {item}")
    if lesson.get("difficult_points"):
        lines.append("")
        lines.append("## 教学难点")
        for item in lesson["difficult_points"]:
            lines.append(f"- {item}")
    if lesson.get("preparation"):
        lines.append("")
        lines.append("## 教学准备")
        for item in lesson["preparation"]:
            lines.append(f"- {item}")
    lines.append("")
    lines.append("## 教学过程")
    lines.append("")
    lines.append("| 环节 | 时间 | 教师活动 | 学生活动 | 教学意图 |")
    lines.append("|---|---|---|---|---|")
    for stage in lesson.get("process", []):
        lines.append(
            f"| {stage.get('stage', '')} | {stage.get('minutes', '')} 分钟 | {stage.get('teacher_activity', '')} "
            f"| {stage.get('student_activity', '')} | {stage.get('intent', '')} |"
        )
    if lesson.get("board_design"):
        lines.append("")
        lines.append("## 板书 / 画面设计")
        lines.append(lesson["board_design"])
    if lesson.get("homework"):
        lines.append("")
        lines.append("## 课后任务")
        for item in lesson["homework"]:
            lines.append(f"- {item}")
    return "\n".join(lines)


def _worksheet_markdown(course: str, worksheet: dict[str, Any], *, include_answers: bool, label: str | None) -> str:
    suffix = "（教师版）" if include_answers else "（学生版）"
    lines = [f"# {course} — {worksheet.get('title', '')} 学习单{suffix}"]
    if label:
        lines.append(f"> **{label}**")
    lines.append("")
    lines.append(worksheet.get("instructions", ""))
    for index, task in enumerate(worksheet.get("tasks", []), start=1):
        lines.append("")
        lines.append(f"## 任务 {index}：{task.get('prompt', '')}")
        if task.get("fields"):
            lines.append("| " + " | ".join(task["fields"]) + " |")
            lines.append("|" + "---|" * len(task["fields"]))
            lines.append("|" + "  |" * len(task["fields"]))
        if include_answers and task.get("answer"):
            lines.append(f"**答案：** {task['answer']}")
            if task.get("answer_notes"):
                lines.append(f"**解析：** {task['answer_notes']}")
    if worksheet.get("reflection_questions"):
        lines.append("")
        lines.append("## 反思问题")
        for question in worksheet["reflection_questions"]:
            lines.append(f"- {question}")
    if not include_answers:
        lines.append("")
        lines.append(f"*{worksheet.get('student_version_note', '学生版不含答案')}*")
    return "\n".join(lines)


def _write_text_docx(path: Path, title: str, lines: list[str]) -> None:
    import docx

    document = docx.Document()
    _set_cjk_font(document)
    document.add_heading(title, level=1)
    for line in lines:
        text = line.strip()
        if not text:
            continue
        if text.startswith("#"):
            level = min(len(text) - len(text.lstrip("#")), 3)
            document.add_heading(text.lstrip("# ").strip(), level=level)
            continue
        document.add_paragraph(text)
    document.save(str(path))


def write_exports(ctx: Any, *, mode: str, gate_ctx: gates.GateContext, note: str = "") -> dict[str, Any]:
    requirements = ctx.content("requirements") or {}
    course = requirements.get("course", "课程")
    readiness = gates.export_readiness_with_overrides(gate_ctx)
    blocking = list(readiness.get("blocking", []))
    if note:
        blocking.insert(0, note)
    label = LABEL_FORMAL if mode == "formal" else LABEL_DRAFT
    output_dir = _output_dir(ctx, mode)
    output_dir.mkdir(parents=True, exist_ok=True)
    spec = ctx.content("template_spec") or {}
    columns = [c["name"] for c in spec.get("columns", [])]
    files: list[dict[str, Any]] = []
    sections = (requirements.get("sections") or [])
    for section in sections:
        artifact_id = f"script:{section['section_id']}"
        script = ctx.content(artifact_id) or {}
        rev_id = ctx.rev(artifact_id) or "r0"
        version = rev_id.split("@")[-1] if "@" in rev_id else rev_id
        stem = f"{_safe(course)}_{section['section_id']}.{version}"
        md_path = output_dir / f"{stem}.md"
        md_path.write_text(
            _section_markdown(course, script, LABEL_DRAFT if mode == "draft" else None, blocking if mode == "draft" else []),
            encoding="utf-8",
        )
        files.append(
            {
                "path": str(md_path),
                "kind": "markdown",
                "sha256": file_hash(md_path),
                "label": label,
                "role": "lesson_script",
            }
        )
        docx_path = output_dir / f"{stem}.docx"
        _write_docx(
            docx_path,
            course,
            script,
            columns,
            LABEL_DRAFT if mode == "draft" else None,
            blocking if mode == "draft" else [],
        )
        files.append(
            {
                "path": str(docx_path),
                "kind": "docx",
                "sha256": file_hash(docx_path),
                "label": label,
                "role": "lesson_script",
            }
        )

        def _add_media(source_path: str, kind: str, role: str) -> None:
            source = Path(source_path)
            if not source.exists():
                return
            target = output_dir / f"{_safe(course)}_{section['section_id']}_{source.name}"
            shutil.copyfile(source, target)
            files.append(
                {"path": str(target), "kind": kind, "sha256": file_hash(target), "label": label, "role": role}
            )

        diagrams_doc = ctx.content(f"diagrams:{section['section_id']}") or {}
        for diagram in diagrams_doc.get("diagrams", []):
            _add_media(str(diagram.get("svg_path", "")), "svg", "diagram")
        storyboard_doc = ctx.content(f"storyboard:{section['section_id']}") or {}
        if storyboard_doc.get("shots"):
            media_base = Path(ctx.store.base_dir) / "media"
            _add_media(str(media_base / "storyboard" / f"{section['section_id']}.json"), "json", "storyboard")
            _add_media(str(media_base / "storyboard" / f"{section['section_id']}.md"), "markdown", "storyboard")
            _add_media(str(media_base / "prompts" / f"{section['section_id']}.txt"), "txt", "prompt")

        draft_label = LABEL_DRAFT if mode == "draft" else None
        lesson = ctx.content(f"lesson_plan:{section['section_id']}") or {}
        if lesson:
            lesson_path = output_dir / f"{stem}.lesson.md"
            lesson_path.write_text(_lesson_markdown(course, lesson, draft_label), encoding="utf-8")
            files.append(
                {"path": str(lesson_path), "kind": "markdown", "sha256": file_hash(lesson_path), "label": label, "role": "lesson_plan"}
            )
            lesson_docx = output_dir / f"{stem}.lesson.docx"
            _write_text_docx(
                lesson_docx,
                f"{course} — {lesson.get('title', '')}（教案）",
                _lesson_markdown(course, lesson, draft_label).splitlines(),
            )
            files.append(
                {"path": str(lesson_docx), "kind": "docx", "sha256": file_hash(lesson_docx), "label": label, "role": "lesson_plan"}
            )
        worksheet = ctx.content(f"worksheet:{section['section_id']}") or {}
        if worksheet:
            student_path = output_dir / f"{stem}.worksheet.student.md"
            student_path.write_text(
                _worksheet_markdown(course, worksheet, include_answers=False, label=draft_label), encoding="utf-8"
            )
            files.append(
                {"path": str(student_path), "kind": "markdown", "sha256": file_hash(student_path), "label": label, "role": "worksheet"}
            )
            student_docx = output_dir / f"{stem}.worksheet.student.docx"
            _write_text_docx(
                student_docx,
                f"{course} — {worksheet.get('title', '')} 学习单（学生版）",
                _worksheet_markdown(course, worksheet, include_answers=False, label=draft_label).splitlines(),
            )
            files.append(
                {"path": str(student_docx), "kind": "docx", "sha256": file_hash(student_docx), "label": label, "role": "worksheet"}
            )
            teacher_path = output_dir / f"{stem}.worksheet.teacher.md"
            teacher_path.write_text(
                _worksheet_markdown(course, worksheet, include_answers=True, label=draft_label), encoding="utf-8"
            )
            files.append(
                {"path": str(teacher_path), "kind": "markdown", "sha256": file_hash(teacher_path), "label": label, "role": "worksheet_teacher"}
            )


    deck = ctx.content("pptx_deck") or {}
    deck_path = Path(str(deck.get("path", ""))) if deck.get("path") else None
    if deck_path is not None and deck_path.is_file():
        target = output_dir / f"{_safe(course)}_slides.pptx"
        shutil.copyfile(deck_path, target)
        files.append(
            {"path": str(target), "kind": "pptx", "sha256": file_hash(target), "label": label, "role": "presentation"}
        )

    manifest: dict[str, Any] = {
        "project_id": ctx.project_id,
        "mode": mode,
        "draft_label": LABEL_DRAFT if mode == "draft" else None,
        "files": files,
        "gate_summary": readiness.get("summary", []),
        "renderer": _renderer_info(),
        "signoffs": current_signoffs(ctx.store),
        "created_at": now_iso(),
    }
    return manifest


def _renderer_info() -> dict[str, Any]:
    try:
        from importlib.metadata import version as _pkg_version

        docx_version = _pkg_version("python-docx")
    except Exception:  # noqa: BLE001 - informational metadata only
        docx_version = "unknown"
    return {"docx": f"python-docx {docx_version}", "fonts": ["Microsoft YaHei"], "qa": "python-docx 打开/解析检查"}


def _safe(text: str) -> str:
    return "".join(ch if ch.isalnum() or ch in "-_()" or "\u4e00" <= ch <= "\u9fff" else "_" for ch in text)[:40]
